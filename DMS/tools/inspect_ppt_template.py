import sys
from pptx import Presentation

default_path = r"d:\POC\DMS2\DMS\docs\AI_DMS.v0.1.0.pptx"
path = sys.argv[1] if len(sys.argv) > 1 else default_path
p = Presentation(path)
print("slides", len(p.slides))
for i, s in enumerate(p.slides, 1):
    title = ""
    if s.shapes.title is not None and s.shapes.title.has_text_frame:
        title = s.shapes.title.text
    texts = []
    for sh in s.shapes:
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            txt = " ".join([par.text.strip() for par in sh.text_frame.paragraphs if par.text.strip()])
            if txt:
                texts.append(txt[:140])
    print(f"{i:02d}: title={title!r}; text_shapes={len(texts)}")
    for j, t in enumerate(texts[:6], 1):
        print(f"    {j}. {t}")
