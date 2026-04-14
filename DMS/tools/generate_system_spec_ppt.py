from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt


TITLE = "DMS Solution System Specification"
SUBTITLE = "Document Management Solution | April 2026"


def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullet_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
            p.text = text
            p.level = level
        else:
            p.text = str(item)
            p.level = 0
        p.font.size = Pt(20 if p.level == 0 else 16)


def add_two_col_slide(prs, title, left_title, left_bullets, right_title, right_bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title

    left_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(5.9), Inches(5.2))
    right_box = slide.shapes.add_textbox(Inches(6.7), Inches(1.5), Inches(5.9), Inches(5.2))

    left_tf = left_box.text_frame
    left_tf.text = left_title
    left_tf.paragraphs[0].font.bold = True
    left_tf.paragraphs[0].font.size = Pt(22)
    for item in left_bullets:
        p = left_tf.add_paragraph()
        if isinstance(item, tuple):
            txt, level = item
            p.text = txt
            p.level = level
        else:
            p.text = str(item)
            p.level = 0
        p.font.size = Pt(18 if p.level == 0 else 14)

    right_tf = right_box.text_frame
    right_tf.text = right_title
    right_tf.paragraphs[0].font.bold = True
    right_tf.paragraphs[0].font.size = Pt(22)
    for item in right_bullets:
        p = right_tf.add_paragraph()
        if isinstance(item, tuple):
            txt, level = item
            p.text = txt
            p.level = level
        else:
            p.text = str(item)
            p.level = 0
        p.font.size = Pt(18 if p.level == 0 else 14)


def build_presentation(output_path: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs, TITLE, SUBTITLE)

    add_bullet_slide(prs, "1. Solution Overview", [
        "Purpose: Ingest, organize, retrieve, and govern business documents",
        "Primary Users: System administrators, document administrators, contributors, viewers",
        "Core Domains: Document Management, Knowledge Collaboration, System Administration",
        "Architecture Style: Web client + REST backend + data/search services",
    ])

    add_two_col_slide(
        prs,
        "2. Logical Architecture",
        "Frontend Layer",
        [
            "React 19 + Vite single-page application",
            "Function modules: Dashboard, Document, Knowledge, Admin, Auditing, Reports",
            "Upload and Scan workflows with metadata-driven forms",
            "Proxies API requests to backend during development",
        ],
        "Backend and Data Layer",
        [
            "Spring Boot REST API (Java 21)",
            "MySQL for system and document metadata",
            "OpenSearch for indexing and retrieval",
            "Optional integrations: DeepSeek, OCR/extraction services",
        ],
    )

    add_bullet_slide(prs, "3. Major Functional Capabilities", [
        "Document lifecycle: Upload, versioning, metadata update, archive, download",
        "Folder and permission model with role/group-based access controls",
        "Knowledge Collaboration topics with contributions, links, shares, attachments",
        "System Administration: retention, reminders, job scheduling, code tables",
        "Audit and reporting modules for operational traceability",
    ])

    add_bullet_slide(prs, "4. API and Integration Specification", [
        "Document APIs: /api/documents, /versions, /download, /folders/tree",
        "Knowledge APIs: /api/knowledge/topics, /join, /contributions, /share, /chain/download",
        "AI APIs: /api/chatbot/search, /summary, /qa with optional streaming endpoints",
        "Scanner Bridge API (workstation local): /health, /scanners, /scan",
    ])

    add_bullet_slide(prs, "5. Security and Access Control", [
        "Role-based function access managed in application context and backend rules",
        "Folder-level permissions for Read, Write, Delete operations",
        "Operational recommendation: secure secrets via environment variables",
        "CORS policy controlled with APP_CORS_ALLOWED_ORIGINS",
        "Service hardening: keep scanner bridge on localhost only",
    ])

    add_bullet_slide(prs, "6. Environment and Prerequisites", [
        "Client/Web: Modern browser for frontend UI",
        "Frontend runtime: Node.js 20.19+ and npm 10+",
        "Backend runtime: Java 21 and Maven 3.9+",
        "Data stores: MySQL 8+ and OpenSearch endpoint",
        "Optional local services: DeepSeek endpoint, OCR/extraction endpoint",
    ])

    add_bullet_slide(prs, "7. Deployment Topology", [
        "Tier 1: Browser client for user interaction",
        "Tier 2: Spring Boot API service on default port 8080",
        "Tier 3: MySQL and OpenSearch infrastructure services",
        "Client-side local utility: Node Scanner Bridge on port 8787",
        "Recommended startup order: MySQL -> OpenSearch -> Backend -> Frontend",
    ])

    add_bullet_slide(prs, "8. Scanner Bridge Service Specification", [
        "Component: DMS Node Scanner Bridge (Node.js)",
        "Workstation deployment path: DMS/tools/node-scanner-bridge",
        "Windows Service name: DMSNodeScannerBridge",
        "Auto-start mode: automatic delayed start + restart on failure",
        "Installer tooling: install-service.ps1 and uninstall-service.ps1",
        "Health validation: http://localhost:8787/health and /scanners",
    ])

    add_bullet_slide(prs, "9. Operational Administration Specification", [
        "User and group administration with role assignment",
        "Retention and reminder rule governance by category/date basis",
        "Job schedule management for background processing",
        "Code table governance for reusable business values",
        "Backup and recovery planning for MySQL and OpenSearch",
    ])

    add_bullet_slide(prs, "10. Non-Functional Requirements", [
        "Reliability: Service restart policies and health endpoint checks",
        "Maintainability: Modular frontend components and REST service boundaries",
        "Scalability: OpenSearch indexing and paginated APIs",
        "Observability: Logs, audit events, and operational status feedback in UI",
        "Compatibility: Windows workstation support for scanner bridge",
    ])

    add_bullet_slide(prs, "11. Validation and Acceptance Criteria", [
        "System boots with all tiers reachable (frontend, backend, DB, OpenSearch)",
        "Document upload/version workflows execute with correct permission enforcement",
        "Knowledge topic create/share/link/upload actions function end-to-end",
        "Scanner detection and scan operations succeed through local bridge service",
        "Admin modules apply retention/reminder/job/code-table updates successfully",
    ])

    add_bullet_slide(prs, "12. Appendix: Key Runtime Endpoints", [
        "Frontend: http://localhost:5173",
        "Backend: http://localhost:8080",
        "Scanner Bridge: http://localhost:8787",
        "OpenSearch (default): https://localhost:9200",
        "Primary docs: DMS/README.md and frontend/public/system-admin-manual.html",
    ])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))


if __name__ == "__main__":
    out = Path("d:/POC/DMS2/DMS/docs/DMS_System_Specification_Apr2026.pptx")
    build_presentation(out)
    print(f"Generated: {out}")
