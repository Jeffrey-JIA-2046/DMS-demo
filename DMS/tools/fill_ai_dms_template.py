from pathlib import Path
from pptx import Presentation

TEMPLATE = Path(r"d:\POC\DMS2\DMS\docs\AI_DMS.v0.1.0.pptx")
OUTPUT = Path(r"d:\POC\DMS2\DMS\docs\AI_DMS.v0.1.0_filled_system_spec.pptx")


def get_text_shapes(slide):
    return [sh for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.has_text_frame]


def set_shape_text(shape, text):
    tf = shape.text_frame
    tf.clear()
    lines = text.split("\n")
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.text = line
        first = False


def fill():
    prs = Presentation(str(TEMPLATE))

    content = {
        1: [
            "Document Management Solution (DMS)\nSystem Specification\nVersion: April 2026\nPrepared for: System Administration and Technical Teams"
        ],
        2: [
            "Contents\n1. Solution Overview\n2. Major Features\n3. Architecture and Components\n4. Deployment Specification\n5. Licensing and Operations\n6. Q&A"
        ],
        3: [
            "Major Features",
            "Document Management\n- Upload, versioning, metadata update, archive, and secure download\n\nKnowledge Collaboration\n- Topic creation, contribution timeline, document links, sharing, and chain export\n\nSystem Administration\n- Retention, reminder, job scheduling, code tables, and user/group access control"
        ],
        4: [
            "Folder Level Access Rights",
            "Role and group based authorization by folder\n- Read, Write, Delete controls per group\n- Inherit permissions from parent folder\n- Folder specific metadata template support\n- Effective permission filtering in UI and API"
        ],
        5: [
            "Document Upload and Scan",
            "Upload Workflow\n- Upload files with metadata and tags into selected folder\n- Approver and supervisor assignment\n\nScan Workflow\n- Scan from client workstation via local Scanner Bridge\n- Supports scanner discovery and scan-to-PDF/PNG\n- Optional OCR and data extraction after upload"
        ],
        6: [
            "Document Online Preview",
            "In-browser document viewing and version inspection\n- PDF page-by-page preview\n- Version history access and download\n- Metadata side-panel and update workflow\n- Latest and historical file retrieval endpoints"
        ],
        7: [
            "AI Assistant",
            "AI-enabled productivity features\n- Natural language search for related documents\n- Document summary generation\n- Document question and answer\n- Streaming response support for concierge style interaction"
        ],
        8: [
            "AI-powered Search Capability",
            "Search modes\n- Metadata and indexed content search\n- Keyword and filtered queries (owner/category/status/tags/folder)\n- Natural language query support via AI endpoint\n\nData Sources\n- MySQL metadata and OpenSearch index"
        ],
        9: [
            "OCR and Data Extraction",
            "OCR and extraction pipeline\n- Full-text OCR for scanned documents\n- Optional metadata extraction from OCR output\n- Confidence-based extraction validation\n- Integration points configurable by environment variables"
        ],
        10: [
            "Document Collaboration",
            "Collaboration capabilities\n- Multi-user topic participation and contribution\n- Share topics and attach supporting documents\n- Link existing DMS files with context notes\n- Download full knowledge chain for audit/compliance"
        ],
        11: [
            "System Auditing",
            "Governance and traceability\n- Role-based feature controls\n- Transaction-level audit logs\n- Searchable audit records by user/action/time\n- Reporting support for operational review"
        ],
        12: [
            "Deployment",
            "On-Premises Deployment\n- Frontend: React + Vite\n- Backend: Spring Boot (Java 21)\n- Data services: MySQL 8+ and OpenSearch\n- Scanner Bridge as local Windows service on workstation (port 8787)",
            "Hybrid Cloud Deployment\n- Container or VM based hosting for frontend/backend\n- Managed DB and search options supported\n- Scanner Bridge remains on client workstation for local device access\n- API and CORS settings configurable per environment"
        ],
        13: [
            "Licensing Scheme",
            "Core Platform Licensing\n- Application licensing for DMS modules (document, knowledge, admin, audit, reports)\n- Includes API-based integrations and standard support scope",
            "Operational Components\n- Infrastructure and runtime licensing (Java/Node/DB/search/cloud) based on deployment model\n- Optional AI/OCR services licensed per usage or subscription",
            "Recommended Commercial Model\n- Annual subscription with maintenance and enhancement support\n- Optional implementation package for rollout, migration, and training"
        ],
        14: ["Q and A"],
        15: ["Thank You\nDocument Management Solution System Specification"]
    }

    for slide_index, shape_texts in content.items():
        slide = prs.slides[slide_index - 1]
        text_shapes = get_text_shapes(slide)
        limit = min(len(text_shapes), len(shape_texts))
        for i in range(limit):
            set_shape_text(text_shapes[i], shape_texts[i])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    fill()
